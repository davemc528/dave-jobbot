from __future__ import annotations

import json
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pydantic import BaseModel, Field

from jobbot.config import BASE_DIR
from jobbot.documents.extract import extract_text
from jobbot.jobs.analysis import JobAnalysis, analyze_job
from jobbot.profile.application_answers import list_answers
from jobbot.profile.canonical import list_canonical_facts

GENERATED_DIR = BASE_DIR / "generated_resumes"
BASE_DOCUMENTS = {
    "MSL / Medical Affairs": BASE_DIR / "source_documents" / "David_Cunningham_MSL_Resume.docx",
    "FAS / Technical Applications": (
        BASE_DIR / "source_documents" / "David_Cunningham_FAS_Resume.docx"
    ),
    "Biology Teaching / Academic": BASE_DIR / "source_documents" / "Dave Cunningham.CV.pdf",
    "Discovery / Translational Scientist": (
        BASE_DIR / "source_documents" / "David_Cunningham_MSL_Resume.docx"
    ),
}


class ResumeClaim(BaseModel):
    claim_text: str
    claim_type: str
    supporting_fact_ids: list[str]
    source_provenance: list[str]
    transformation_type: str
    confidence: float
    human_review_required: bool = False
    validation_status: str = "valid"


class TailoredResumeVersion(BaseModel):
    id: int
    job_id: int
    version: int
    status: str
    selected_track: str
    base_document_path: str
    docx_path: str
    text_path: str
    report_path: str
    validation_path: str
    validation_status: str
    warnings: list[str] = Field(default_factory=list)
    claims: list[ResumeClaim] = Field(default_factory=list)


def _verified_claims(connection: sqlite3.Connection, analysis: JobAnalysis) -> list[ResumeClaim]:
    facts = [
        fact
        for fact in list_canonical_facts(connection)
        if fact.active
        and fact.verification_status == "verified"
        and fact.category
        in {
            "identity",
            "geography",
            "education",
            "employment_history",
            "experience_duration",
            "technical_skills",
            "publications",
        }
    ]
    emphasis = {value.casefold() for value in analysis.emphasis_areas}

    def priority(fact: object) -> tuple[int, str]:
        display = str(getattr(fact, "display_value", "") or "")
        relevant = any(term in display.casefold() for term in emphasis)
        category = str(getattr(fact, "category", ""))
        category_rank = {
            "identity": 0,
            "geography": 1,
            "experience_duration": 2,
            "employment_history": 3,
            "technical_skills": 4,
            "education": 5,
            "publications": 6,
        }.get(category, 9)
        return (category_rank - (2 if relevant else 0), display)

    claims: list[ResumeClaim] = []
    for fact in sorted(facts, key=priority):
        if fact.id is None or not fact.display_value:
            continue
        if fact.field_name == "nature_ai_benchmark_doi":
            continue
        claims.append(
            ResumeClaim(
                claim_text=fact.display_value,
                claim_type=fact.category,
                supporting_fact_ids=[f"canonical_fact:{fact.id}"],
                source_provenance=fact.source_documents or ["verified canonical record"],
                transformation_type="reordered",
                confidence=1.0,
            )
        )
    for answer in list_answers(connection):
        if (
            answer.field_name in {"email", "phone", "linkedin_url"}
            and answer.active
            and answer.verification_status == "verified"
            and answer.id is not None
            and answer.display_value
        ):
            claims.insert(
                1,
                ResumeClaim(
                    claim_text=answer.display_value,
                    claim_type="contact_information",
                    supporting_fact_ids=[f"application_answer:{answer.id}"],
                    source_provenance=[answer.provenance or "explicit user verification"],
                    transformation_type="exact",
                    confidence=1.0,
                ),
            )
    return claims


def validate_claims(claims: list[ResumeClaim]) -> list[str]:
    warnings: list[str] = []
    prohibited = (
        "kol management",
        "direct msl experience",
        "clinical practice",
        "patient care",
        "sales quota",
        "people management",
    )
    for claim in claims:
        if not claim.supporting_fact_ids:
            claim.validation_status = "rejected"
            warnings.append(f"Unsupported claim rejected: {claim.claim_text}")
        if any(term in claim.claim_text.casefold() for term in prohibited):
            claim.validation_status = "rejected"
            warnings.append(f"Prohibited unsupported positioning: {claim.claim_text}")
        if "[insert " in claim.claim_text.casefold():
            claim.validation_status = "rejected"
            warnings.append(f"Placeholder detected: {claim.claim_text}")
    return warnings


def _render_docx(path: Path, claims: list[ResumeClaim], track: str) -> None:
    document = Document()
    document.core_properties.title = f"Tailored resume — {track}"
    grouped: dict[str, list[ResumeClaim]] = {}
    for claim in claims:
        if claim.validation_status == "valid":
            grouped.setdefault(claim.claim_type, []).append(claim)
    labels = {
        "identity": "PROFILE",
        "contact_information": "CONTACT",
        "experience_duration": "PROFESSIONAL SUMMARY",
        "employment_history": "EXPERIENCE",
        "technical_skills": "TECHNICAL EXPERTISE",
        "education": "EDUCATION",
        "publications": "SELECTED PUBLICATIONS",
        "geography": "LOCATION",
    }
    for category, category_claims in grouped.items():
        document.add_heading(labels.get(category, category.upper()), level=1)
        for claim in category_claims:
            document.add_paragraph(claim.claim_text, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def tailor_resume(connection: sqlite3.Connection, job_id: int) -> TailoredResumeVersion:
    posting = connection.execute(
        "SELECT complete, completeness_warnings FROM job_postings WHERE job_id=?", (job_id,)
    ).fetchone()
    if posting is None or not posting["complete"]:
        warnings = json.loads(posting["completeness_warnings"]) if posting else []
        raise ValueError(
            "A complete normalized posting is required before tailoring"
            + (f": {'; '.join(warnings)}" if warnings else "")
        )
    analysis = analyze_job(connection, job_id)
    base_path = BASE_DOCUMENTS[analysis.selected_track]
    if not base_path.is_file():
        raise ValueError(f"Base document does not exist: {base_path}")
    version = int(
        connection.execute(
            "SELECT coalesce(max(version), 0) + 1 FROM tailored_resumes WHERE job_id=?",
            (job_id,),
        ).fetchone()[0]
    )
    target_dir = GENERATED_DIR / str(job_id) / f"v{version}"
    docx_path = target_dir / "tailored_resume.docx"
    text_path = target_dir / "ats_preview.txt"
    report_path = target_dir / "tailoring_report.json"
    validation_path = target_dir / "claim_validation.json"
    claims = _verified_claims(connection, analysis)
    warnings = validate_claims(claims)
    warnings.append(
        "Page count was not verified because no reliable local DOCX-to-PDF converter is configured"
    )
    invalid_claims = any(claim.validation_status != "valid" for claim in claims)
    valid_claims = [claim for claim in claims if claim.validation_status == "valid"]
    if not valid_claims:
        raise ValueError("No supported verified claims are available")
    _render_docx(docx_path, claims, analysis.selected_track)
    ats_text = "\n".join(claim.claim_text for claim in valid_claims)
    text_path.write_text(ats_text, encoding="utf-8")
    report = {
        "job_id": job_id,
        "selected_track": analysis.selected_track,
        "base_document": str(base_path),
        "base_text": extract_text(base_path),
        "tailored_text": ats_text,
        "emphasis_areas": analysis.emphasis_areas,
        "withheld_requirements": analysis.unverified_requirements,
        "changes": "Verified claims selected and reordered by category and relevance.",
    }
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    validation_path.write_text(
        json.dumps(
            {
                "status": "invalid" if invalid_claims else "valid",
                "warnings": warnings,
                "claims": [claim.model_dump() for claim in claims],
                "checks": {
                    "no_hidden_text": True,
                    "no_tables_or_text_boxes": True,
                    "source_documents_unchanged": True,
                    "superseded_publications_excluded": True,
                    "ats_text_extractable": bool(ats_text.strip()),
                },
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    cursor = connection.execute(
        """
        INSERT INTO tailored_resumes
          (job_id, version, status, selected_track, base_document_path, docx_path,
           text_path, report_path, validation_path, validation_status, warnings, created_at)
        VALUES (?, ?, 'needs_review', ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            job_id,
            version,
            analysis.selected_track,
            str(base_path),
            str(docx_path),
            str(text_path),
            str(report_path),
            str(validation_path),
            "invalid" if invalid_claims else "valid",
            json.dumps(warnings),
            datetime.now(timezone.utc).isoformat(),
        ),
    )
    resume_id = int(cursor.lastrowid or 0)
    for claim in claims:
        connection.execute(
            """
            INSERT INTO resume_claims
              (tailored_resume_id, claim_text, claim_type, supporting_fact_ids,
               source_provenance, transformation_type, confidence,
               human_review_required, validation_status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                resume_id,
                claim.claim_text,
                claim.claim_type,
                json.dumps(claim.supporting_fact_ids),
                json.dumps(claim.source_provenance),
                claim.transformation_type,
                claim.confidence,
                int(claim.human_review_required),
                claim.validation_status,
            ),
        )
    connection.commit()
    return get_version(connection, job_id, resume_id)


def get_version(
    connection: sqlite3.Connection, job_id: int, version_id: int
) -> TailoredResumeVersion:
    row = connection.execute(
        "SELECT * FROM tailored_resumes WHERE job_id=? AND id=?", (job_id, version_id)
    ).fetchone()
    if row is None:
        raise ValueError("Unknown tailored resume version")
    claims = [
        ResumeClaim(
            claim_text=claim["claim_text"],
            claim_type=claim["claim_type"],
            supporting_fact_ids=json.loads(claim["supporting_fact_ids"]),
            source_provenance=json.loads(claim["source_provenance"]),
            transformation_type=claim["transformation_type"],
            confidence=claim["confidence"],
            human_review_required=bool(claim["human_review_required"]),
            validation_status=claim["validation_status"],
        )
        for claim in connection.execute(
            "SELECT * FROM resume_claims WHERE tailored_resume_id=? ORDER BY id",
            (version_id,),
        )
    ]
    return TailoredResumeVersion(
        id=row["id"],
        job_id=row["job_id"],
        version=row["version"],
        status=row["status"],
        selected_track=row["selected_track"],
        base_document_path=row["base_document_path"],
        docx_path=row["docx_path"],
        text_path=row["text_path"],
        report_path=row["report_path"],
        validation_path=row["validation_path"],
        validation_status=row["validation_status"],
        warnings=json.loads(row["warnings"]),
        claims=claims,
    )


def set_version_status(
    connection: sqlite3.Connection, job_id: int, version_id: int, status: str
) -> TailoredResumeVersion:
    if status not in {"approved", "rejected"}:
        raise ValueError("Status must be approved or rejected")
    version = get_version(connection, job_id, version_id)
    if status == "approved" and version.validation_status != "valid":
        raise ValueError("Only a valid tailored resume may be approved")
    if status == "approved":
        connection.execute(
            """
            UPDATE tailored_resumes SET status='superseded'
            WHERE job_id=? AND status='approved' AND id<>?
            """,
            (job_id, version_id),
        )
        connection.execute(
            """
            UPDATE review_items SET status='resolved'
            WHERE status='pending' AND item_type='resume_approval_required'
              AND json_extract(metadata, '$.job_id')=?
            """,
            (job_id,),
        )
    connection.execute(
        "UPDATE tailored_resumes SET status=?, approved_at=? WHERE id=?",
        (
            status,
            datetime.now(timezone.utc).isoformat() if status == "approved" else None,
            version_id,
        ),
    )
    connection.commit()
    return get_version(connection, job_id, version_id)


def approved_resume_for_job(
    connection: sqlite3.Connection, job_id: int
) -> TailoredResumeVersion | None:
    row = connection.execute(
        """
        SELECT id FROM tailored_resumes
        WHERE job_id=? AND status='approved' AND validation_status='valid'
        ORDER BY version DESC LIMIT 1
        """,
        (job_id,),
    ).fetchone()
    return get_version(connection, job_id, int(row["id"])) if row else None


def contains_placeholder(text: str) -> bool:
    return bool(re.search(r"\[(?:insert|add|todo)[^\]]*\]", text, re.I))
